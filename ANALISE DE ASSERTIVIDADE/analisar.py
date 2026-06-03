#!/usr/bin/env python3
"""
analisar.py  —  Lotinha Analytics: Análise de Cobertura de Predições
━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
Uso:
    python analisar.py                   # pasta atual
    python analisar.py /caminho/pasta    # pasta específica
    python analisar.py . --data 2026-05-21   # apenas uma data
    python analisar.py . --salvar        # salva relatório em .txt

Estrutura esperada:
    pasta/
    ├── PREDIÇÕES/             ← PDFs gerados pelo sistema
    │   └── *21-05-2026*.pdf
    └── RESULTADOS OFICIAIS/   ← XLSXs dos resultados
        └── *21-05-2026*.xlsx
"""

import os, sys, re, subprocess, argparse
from pathlib import Path
from datetime import datetime

# ═══════════════════════════════════════════════════════════════════
#  PARSING DE PDF
# ═══════════════════════════════════════════════════════════════════

def parse_predictions_pdf(path):
    """
    Lê um PDF de predições e retorna:
      (date_str, n_pred, {(banca, hora): frozenset(int)})
    """
    r = subprocess.run(
        ['pdftotext', '-layout', str(path), '-'],
        capture_output=True, text=True
    )
    if r.returncode != 0:
        raise RuntimeError(f"pdftotext falhou em '{path}': {r.stderr.strip()}")
    text = r.stdout

    # Data do sorteio previsto
    date = None
    m = re.search(r'Predições\s*[—–\-]+\s*(\d{4}-\d{2}-\d{2})', text)
    if m:
        date = m.group(1)

    # Quantidade de números preditos por horário
    n_pred = None
    m = re.search(r'N preditos:\s*(\d+)', text)
    if m:
        n_pred = int(m.group(1))

    entries = {}
    lines = text.splitlines()

    for i, line in enumerate(lines):
        # Detecta linha de entrada: "LOTINHA FEDERAL/PONTO  XXh  <numeros>  XX.X%  NNN"
        m = re.match(r'\s*(LOTINHA\s+(?:FEDERAL|PONTO))\s+(\d{1,2}h)\s+(.*)', line)
        if not m:
            continue

        banca = ' '.join(m.group(1).split())   # normaliza espaços
        hora  = m.group(2)
        rest  = m.group(3)

        # Remove confiança% e histórico do fim da linha
        rest_clean = re.sub(r'\s+\d+\.\d+%\s+\d+\s*$', '', rest)

        # Linha de continuação (indentada, com números)
        cont = ''
        if i + 1 < len(lines):
            nxt = lines[i + 1]
            if len(nxt) > 2 and nxt[0] == ' ' and re.search(r'\d', nxt):
                cont = nxt

        all_text = rest_clean + ' ' + cont
        nums = frozenset(
            int(n) for n in re.findall(r'\b(\d{1,2})\b', all_text)
            if 1 <= int(n) <= 25
        )

        if len(nums) >= 10:           # sanidade: deve ter ao menos 10 números
            entries[(banca, hora)] = nums

    return date, n_pred, entries


# ═══════════════════════════════════════════════════════════════════
#  PARSING DE XLSX
# ═══════════════════════════════════════════════════════════════════

def _ensure_openpyxl():
    try:
        import openpyxl
        return openpyxl
    except ImportError:
        subprocess.run(
            [sys.executable, '-m', 'pip', 'install', 'openpyxl', '-q',
             '--break-system-packages'],
            check=True
        )
        import openpyxl
        return openpyxl


def parse_results_xlsx(path):
    """
    Lê um XLSX de resultados e retorna:
      (date_str, {(banca, hora): frozenset(int)})
    """
    openpyxl = _ensure_openpyxl()
    wb = openpyxl.load_workbook(str(path), read_only=True)
    ws = wb.active

    date = None
    m = re.search(r'(\d{4}-\d{2}-\d{2})', wb.sheetnames[0])
    if m:
        date = m.group(1)

    results = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        banca_raw = str(row[0]).strip().upper()
        hora      = str(row[1]).strip() if row[1] else ''
        nums_str  = str(row[2]).strip() if len(row) > 2 and row[2] else ''

        # Pula cabeçalho e linhas vazias
        if not banca_raw or 'SORTEIO' in banca_raw or 'BANCA' in banca_raw:
            continue
        if not nums_str or not re.search(r'\d', nums_str):
            continue

        # Normaliza banca: remove sufixo de hora se presente
        # ex.: "LOTINHA PONTO 07H" → "LOTINHA PONTO"
        banca = re.sub(r'\s+\d{1,2}H\s*$', '', banca_raw).strip()

        nums = frozenset(
            int(n) for n in re.findall(r'\b(\d{1,2})\b', nums_str)
            if 1 <= int(n) <= 25
        )
        if len(nums) == 15:
            results[(banca, hora)] = nums

    return date, results


# ═══════════════════════════════════════════════════════════════════
#  ANÁLISE
# ═══════════════════════════════════════════════════════════════════

def analyze(predictions, results):
    """
    Para cada resultado, verifica:
      - se a predição do mesmo horário cobre todas as 15 dezenas
      - quais outras predições (mesmo banca, outro horário) cobrem
      - melhor sobreposição parcial quando nenhuma cobre completamente
    Retorna lista de dicts.
    """
    rows = []

    for (banca, hora) in sorted(results.keys(), key=lambda x: (x[0], x[1])):
        result_nums = results[(banca, hora)]

        same_covered   = False
        same_overlap   = 0
        cross_covers   = []       # horas cruzadas que cobrem completamente
        best_overlap   = 0
        best_hora      = None

        # Apenas predições da mesma banca
        same_banca_preds = {k: v for k, v in predictions.items() if k[0] == banca}

        # Verifica mesmo horário
        if (banca, hora) in predictions:
            pred = predictions[(banca, hora)]
            overlap = len(result_nums & pred)
            same_covered = result_nums.issubset(pred)
            same_overlap = overlap
            if overlap > best_overlap:
                best_overlap = overlap
                best_hora = hora

        # Verifica outros horários
        for (pb, ph), pred_nums in same_banca_preds.items():
            if ph == hora:
                continue
            overlap = len(result_nums & pred_nums)
            if result_nums.issubset(pred_nums):
                cross_covers.append(ph)
            if overlap > best_overlap:
                best_overlap = overlap
                best_hora = ph

        rows.append({
            'banca':        banca,
            'hora':         hora,
            'result':       result_nums,
            'same_covered': same_covered,
            'same_overlap': same_overlap,
            'cross_covers': sorted(cross_covers),
            'best_overlap': best_overlap,
            'best_hora':    best_hora,
            'covered_any':  same_covered or bool(cross_covers),
        })

    return rows


def compute_pred_efficiency(predictions, results, rows):
    """Para cada predição, quantos resultados ela cobre (total)."""
    hits = {k: 0 for k in predictions}
    for row in rows:
        banca = row['banca']
        res   = row['result']
        for (pb, ph), pred_nums in predictions.items():
            if pb == banca and res.issubset(pred_nums):
                hits[(pb, ph)] += 1
    return hits


# ═══════════════════════════════════════════════════════════════════
#  RELATÓRIO
# ═══════════════════════════════════════════════════════════════════

W = 64   # largura do relatório

def fmt(nums):
    return '  '.join(f'{n:02d}' for n in sorted(nums))

def pct(a, b):
    return f'{a/b*100:.1f}%' if b else '0.0%'

def bar(n, total, width=20):
    filled = round(n / total * width) if total else 0
    return '█' * filled + '░' * (width - filled)


def build_report(date, n_pred, predictions, results, rows, pred_hits, filename_pred, filename_result):
    out = []
    A = out.append

    total         = len(rows)
    same_hits     = [r for r in rows if r['same_covered']]
    cross_only    = [r for r in rows if not r['same_covered'] and r['cross_covers']]
    covered_any   = [r for r in rows if r['covered_any']]
    misses        = [r for r in rows if not r['covered_any']]

    same_pct_val  = len(same_hits) / total * 100 if total else 0
    any_pct_val   = len(covered_any) / total * 100 if total else 0
    miss_pct_val  = len(misses) / total * 100 if total else 0
    avg_miss_ov   = sum(r['best_overlap'] for r in misses) / len(misses) if misses else 0

    ts = datetime.now().strftime('%d/%m/%Y %H:%M')

    # ── Cabeçalho ─────────────────────────────────────────────────
    A('═' * W)
    A(f'  LOTINHA ANALYTICS  ·  ANÁLISE DE COBERTURA')
    A(f'  Data analisada: {date}   |   N preditos / sorteio: {n_pred}')
    A(f'  Gerado em: {ts}')
    A('═' * W)
    A(f'  Predições:  {filename_pred}')
    A(f'  Resultados: {filename_result}')
    A('')

    # ── Resumo ────────────────────────────────────────────────────
    A('  RESUMO GERAL')
    A(f'  {"─"*58}')
    A(f'  {"Sorteios analisados":<42} {total:>5}')
    A(f'  {"Predições disponíveis":<42} {len(predictions):>5}')
    A(f'  {"─"*58}')
    A(f'  {"✅  Acerto mesmo horário":<42} {len(same_hits):>5}  ({pct(len(same_hits), total)})')
    A(f'  {"🔄  Acerto horário cruzado (adicional)":<42} {len(cross_only):>5}')
    A(f'  {"✅  Cobertos por ao menos 1 predição":<42} {len(covered_any):>5}  ({pct(len(covered_any), total)})')
    A(f'  {"❌  Sem cobertura alguma":<42} {len(misses):>5}  ({pct(len(misses), total)})')
    if misses:
        A(f'  {"   Sobreposição parcial média (erros)":<42} {avg_miss_ov:.1f}/15')
    A('')

    # ── Acertos mesmo horário ────────────────────────────────────
    if same_hits:
        A('─' * W)
        A('  ✅  ACERTOS — MESMO HORÁRIO')
        A('')
        for r in same_hits:
            A(f'  {r["banca"]}  {r["hora"]}')
            A(f'     Dezenas sorteadas:  {fmt(r["result"])}')
            p = predictions.get((r['banca'], r['hora']), set())
            missed_by_pred = sorted(r['result'] - p) if p else []
            A('')

    # ── Acertos cruzados ─────────────────────────────────────────
    if cross_only:
        A('─' * W)
        A('  🔄  ACERTOS CRUZADOS  (palpite de outro horário cobriu o resultado)')
        A('')
        for r in cross_only:
            covers_str = '  '.join(r['cross_covers'])
            A(f'  Resultado {r["hora"]:>4s}  ←  coberto pelo(s) palpite(s):  {covers_str}')
            A(f'     Dezenas: {fmt(r["result"])}')
            A('')

    # ── Erros ────────────────────────────────────────────────────
    A('─' * W)
    A('  ❌  SEM COBERTURA')
    A('')
    for r in misses:
        nota = f'melhor sobreposição: {r["best_overlap"]}/15  (palpite {r["best_hora"]})'
        A(f'  {r["banca"]}  {r["hora"]}   [{nota}]')
        A(f'     Dezenas: {fmt(r["result"])}')
        A('')

    # ── Eficiência por predição ───────────────────────────────────
    A('─' * W)
    A('  📊  EFICIÊNCIA POR PALPITE  (quantos sorteios cada predição cobre)')
    A('')
    max_hits = max(pred_hits.values(), default=1) or 1
    for (pb, ph), hits in sorted(pred_hits.items(), key=lambda x: (-x[1], x[0][1])):
        b = bar(hits, max_hits)
        A(f'  Palpite {ph:>4s}  {hits:2d} sorteio(s)  {b}')
    A('')

    # ── Índice de eficiência ──────────────────────────────────────
    A('─' * W)
    A('  🎯  ÍNDICE DE EFICIÊNCIA DO SISTEMA')
    A('')
    A(f'  Taxa de acerto mesmo horário:        {same_pct_val:6.1f}%  {bar(len(same_hits), total)}')
    A(f'  Taxa de cobertura (qualquer hora):   {any_pct_val:6.1f}%  {bar(len(covered_any), total)}')
    A(f'  Taxa de erro (sem cobertura alguma): {miss_pct_val:6.1f}%  {bar(len(misses), total)}')
    A('')

    if same_pct_val >= 60:
        nivel = '🟢  ALTA  — sistema com excelente precisão (≥ 60%)'
    elif same_pct_val >= 35:
        nivel = '🟡  MÉDIA — há margem para melhoria (35–60%)'
    else:
        nivel = '🔴  BAIXA — sistema precisa de ajustes significativos (< 35%)'
    A(f'  Avaliação: {nivel}')
    A('')
    A('═' * W)
    A('  ⚠️   Aviso: predições baseadas em frequência histórica não garantem')
    A('       acertos futuros. Use apenas para estudo e análise estatística.')
    A('═' * W)

    return '\n'.join(out)


# ═══════════════════════════════════════════════════════════════════
#  UTILITÁRIOS
# ═══════════════════════════════════════════════════════════════════

def date_from_filename(name):
    """Extrai YYYY-MM-DD de nomes como '..._21-05-2026.pdf' ou '..._2026-05-21.xlsx'"""
    # DD-MM-YYYY
    m = re.search(r'(\d{2})-(\d{2})-(\d{4})', name)
    if m:
        return f'{m.group(3)}-{m.group(2)}-{m.group(1)}'
    # YYYY-MM-DD
    m = re.search(r'(\d{4})-(\d{2})-(\d{2})', name)
    if m:
        return m.group(0)
    return None


def check_dependencies():
    """Verifica que pdftotext está disponível."""
    r = subprocess.run(['which', 'pdftotext'], capture_output=True)
    if r.returncode != 0:
        print('❌  pdftotext não encontrado.')
        print('   Instale com:  sudo apt install poppler-utils')
        sys.exit(1)


# ═══════════════════════════════════════════════════════════════════
#  MAIN
# ═══════════════════════════════════════════════════════════════════

def main():
    parser = argparse.ArgumentParser(
        description='Lotinha Analytics — Análise de Cobertura de Predições'
    )
    parser.add_argument(
        'pasta', nargs='?', default='.',
        help='Pasta raiz com PREDIÇÕES/ e RESULTADOS OFICIAIS/ (padrão: .)'
    )
    parser.add_argument(
        '--data', metavar='YYYY-MM-DD',
        help='Analisar apenas esta data (ex: 2026-05-21)'
    )
    parser.add_argument(
        '--salvar', action='store_true',
        help='Salvar relatório em relatorio_YYYY-MM-DD.txt na pasta raiz'
    )
    args = parser.parse_args()

    check_dependencies()

    root        = Path(args.pasta).resolve()
    pred_dir    = root / 'PREDIÇÕES'
    result_dir  = root / 'RESULTADOS OFICIAIS'

    for d, label in [(pred_dir, 'PREDIÇÕES'), (result_dir, 'RESULTADOS OFICIAIS')]:
        if not d.exists():
            print(f"❌  Pasta '{label}' não encontrada em '{root}'")
            print(f"   Crie a pasta e coloque os arquivos corretos nela.")
            sys.exit(1)

    pred_files   = sorted(pred_dir.glob('*.pdf'))
    result_files = sorted(result_dir.glob('*.xlsx')) + sorted(result_dir.glob('*.xls'))

    if not pred_files:
        print(f"❌  Nenhum PDF encontrado em '{pred_dir}'")
        sys.exit(1)
    if not result_files:
        print(f"❌  Nenhum XLSX encontrado em '{result_dir}'")
        sys.exit(1)

    pred_by_date   = {date_from_filename(f.name): f for f in pred_files   if date_from_filename(f.name)}
    result_by_date = {date_from_filename(f.name): f for f in result_files if date_from_filename(f.name)}

    common = sorted(set(pred_by_date) & set(result_by_date))

    if args.data:
        common = [d for d in common if d == args.data]

    if not common:
        print('⚠️   Nenhuma data em comum entre predições e resultados.')
        print(f'   Predições:  {sorted(pred_by_date.keys())}')
        print(f'   Resultados: {sorted(result_by_date.keys())}')
        sys.exit(1)

    for date in common:
        pf = pred_by_date[date]
        rf = result_by_date[date]

        print(f'\n🔍  Analisando {date}...', file=sys.stderr)

        try:
            pred_date, n_pred, predictions = parse_predictions_pdf(pf)
        except Exception as e:
            print(f'❌  Erro ao ler predição {pf.name}: {e}', file=sys.stderr)
            continue

        try:
            result_date, results = parse_results_xlsx(rf)
        except Exception as e:
            print(f'❌  Erro ao ler resultado {rf.name}: {e}', file=sys.stderr)
            continue

        if not predictions:
            print(f'⚠️   Nenhuma predição extraída de {pf.name}', file=sys.stderr)
            continue
        if not results:
            print(f'⚠️   Nenhum resultado extraído de {rf.name}', file=sys.stderr)
            continue

        rows      = analyze(predictions, results)
        pred_hits = compute_pred_efficiency(predictions, results, rows)
        report    = build_report(
            date, n_pred, predictions, results, rows, pred_hits,
            pf.name, rf.name
        )

        print(report)

        if args.salvar:
            out_path = root / f'relatorio_{date}.txt'
            out_path.write_text(report, encoding='utf-8')
            print(f'\n📄  Relatório salvo em: {out_path}', file=sys.stderr)

    if len(common) > 1:
        print(f'\n{"─"*64}', file=sys.stderr)
        print(f'✅  {len(common)} data(s) analisada(s): {", ".join(common)}', file=sys.stderr)


if __name__ == '__main__':
    main()
